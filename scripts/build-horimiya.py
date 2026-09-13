#!/usr/bin/env python3
"""ホリミヤ構文リストを Excel / CSV / JS に変換する。

source of truth は data/horimiya.xlsx（無ければ data/horimiya.tsv）。
列は 時間 / 構文 / 例文 / 日本語 の4列で、2列目と4列目をアプリに取り込む。

    python3 scripts/build-horimiya.py              # xlsx優先
    python3 scripts/build-horimiya.py --from tsv   # tsvを強制
"""

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
TSV = ROOT / 'data' / 'horimiya.tsv'
XLSX = ROOT / 'data' / 'horimiya.xlsx'
CSV_OUT = ROOT / 'data' / 'horimiya-grammar.csv'
JS_OUT = ROOT / 'words-horimiya-grammar.js'

HEADERS = ['時間', '構文', '例文', '日本語']


def read_tsv():
    rows = []
    with TSV.open(encoding='utf-8') as f:
        for line in f:
            line = line.rstrip('\n').rstrip('\r')
            if not line.strip():
                continue
            cells = line.split('\t')
            cells += [''] * (4 - len(cells))
            rows.append([c.strip() for c in cells[:4]])
    return rows


def read_xlsx():
    wb = load_workbook(XLSX)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = ['' if c is None else str(c).strip() for c in row]
        cells += [''] * (4 - len(cells))
        if not any(cells[:4]):
            continue
        rows.append(cells[:4])
    return rows


def drop_header(rows):
    if rows and rows[0][1] in ('構文', '英語', 'English'):
        return rows[1:]
    return rows


def to_words(rows):
    words = []
    seen = set()
    for time, pattern, example, jp in rows:
        if not pattern or not jp:
            continue
        if pattern in seen:
            print(f'  重複スキップ: {pattern}')
            continue
        seen.add(pattern)
        words.append({'de': pattern, 'jp': jp, 'time': time, 'example': example})
    return words


def write_xlsx(words):
    wb = Workbook()
    ws = wb.active
    ws.title = 'ホリミヤ'
    ws.append(HEADERS)
    for w in words:
        ws.append([w['time'], w['de'], w['example'], w['jp']])
    widths = [10, 44, 52, 46]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = 'A2'
    wb.save(XLSX)


def write_tsv(words):
    lines = ['\t'.join(HEADERS)]
    for w in words:
        lines.append('\t'.join([w['time'], w['de'], w['example'], w['jp']]))
    TSV.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def write_csv(words):
    with CSV_OUT.open('w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['英語', '日本語'])
        for w in words:
            writer.writerow([w['de'], w['jp']])
    risky = [w['de'] for w in words if ',' in w['de']]
    if risky:
        print('  ⚠ 英語側にカンマを含む行があります（CSVフォールバック時に崩れます）:')
        for r in risky:
            print(f'    {r}')


def write_js(words):
    lines = ['window.HORIMIYA_GRAMMAR_WORDS = [']
    for w in words:
        de = json.dumps(w['de'], ensure_ascii=False)
        jp = json.dumps(w['jp'], ensure_ascii=False)
        lines.append(f'  {{de:{de}, jp:{jp}}},')
    lines.append('];')
    JS_OUT.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--from', dest='source', choices=['xlsx', 'tsv'], default=None)
    args = parser.parse_args()

    source = args.source or ('xlsx' if XLSX.exists() else 'tsv')
    rows = drop_header(read_xlsx() if source == 'xlsx' else read_tsv())
    words = to_words(rows)
    if not words:
        raise SystemExit('構文が0件です。元ファイルを確認してください。')

    write_xlsx(words)
    write_tsv(words)
    write_csv(words)
    write_js(words)

    levels = (len(words) + 19) // 20
    print(f'source: {source}')
    print(f'{len(words)}構文 → Level {levels} 個分')
    for path in (XLSX, TSV, CSV_OUT, JS_OUT):
        print(f'  wrote {path.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
